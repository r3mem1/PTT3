"""Деревянная (без обращения к ИИ) проверка рендера всех layout'ов.

Строит один синтетический PresentationData, покрывающий все существующие
type/layout комбинации (включая пути даунгрейда chart/comparison при
нехватке данных), и напрямую вызывает pptx_builder.build_presentation —
без Telegram и без реального вызова OpenRouter. Быстро ловит геометрические
регрессии (наложение текста, переполнение слайда) ценой того, что не может
поймать регрессии в самом промпте (для этого — агент prompt-tester).

Запуск: python scripts/render_check.py (из корня репозитория bot/).
"""

from __future__ import annotations

import asyncio
import shutil
import sys
import tempfile
from pathlib import Path

BOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BOT_DIR))

from PIL import Image  # noqa: E402

from models import normalize_presentation  # noqa: E402
from pptx_builder import build_presentation  # noqa: E402

RAW = {
    "title": "Test Deck — All Layouts",
    "slides": [
        {"type": "title", "title": "Все раскладки на одном прогоне", "subtitle": "render_check.py"},
        {"type": "content", "layout": "text", "title": "Обычный текст (буллеты)", "bullets": [
            "Стоит отметить, что это буллет с водой в начале — должна вырезаться.",
            "**Жирный** буллет с протёкшим markdown — должен вычиститься.",
            "- пункт с лишним дефисом впереди",
            "Очень длинное предложение о том, как важно тестировать рендеринг. Второе предложение должно быть отрезано.",
        ]},
        {"type": "content", "layout": "text", "title": "Слитный текст (body)", "body": (
            "Стоит отметить, что это связный абзац вместо буллетов. Модель иногда решает, "
            "что мысль лучше читается слитно. Третье предложение продолжает мысль."
        )},
        {"type": "content", "layout": "image_left", "title": "Картинка слева (буллеты)", "bullets": ["Пункт раз", "Пункт два"], "image_query": "office team"},
        {"type": "content", "layout": "image_left", "title": "Картинка слева (body)", "body": "Связный абзац рядом с картинкой вместо списка.", "image_query": "reading book"},
        {"type": "content", "layout": "image_right", "title": "Картинка справа", "bullets": ["Пункт раз", "Пункт два"], "image_query": "data analytics"},
        {"type": "content", "layout": "cards", "title": "Карточки", "cards": [
            {"title": "Шаг 1", "body": "Описание первого шага"},
            {"title": "Шаг 2", "body": "Описание второго шага"},
            {"title": "Шаг 3", "body": "Описание третьего шага"},
        ]},
        {"type": "content", "layout": "stat", "title": "Метрика", "metric": "37%", "metric_label": "Рост выручки", "bullets": ["Пояснение к цифре и источник"]},
        {"type": "content", "layout": "chart", "title": "Bar chart", "chart": {
            "type": "bar", "categories": ["2023", "2024", "2025"],
            "series": [{"name": "Выручка", "values": [12, 18, 27]}, {"name": "Расходы", "values": [8, 10, 14]}],
        }},
        {"type": "content", "layout": "chart", "title": "Pie chart", "chart": {
            "type": "pie", "categories": ["A", "B", "C", "D"],
            "series": [{"name": "Доля", "values": [40, 25, 20, 15]}],
        }},
        {"type": "content", "layout": "chart", "title": "Line chart", "chart": {
            "type": "line", "categories": ["Янв", "Фев", "Мар", "Апр"],
            "series": [{"name": "Пользователи", "values": [100, 140, 190, 260]}],
        }},
        {"type": "content", "layout": "chart", "title": "Недостаточно данных -> должен стать text", "chart": {
            "type": "bar", "categories": ["Только одна"], "series": [{"name": "X", "values": [1]}],
        }, "bullets": ["Фолбэк на текст, раз данных мало"]},
        {"type": "content", "layout": "timeline", "title": "Дорожная карта", "timeline": [
            {"label": "Q1 2026", "title": "Запуск MVP", "body": "Первая версия продукта"},
            {"label": "Q2 2026", "title": "Пилот", "body": "Тестирование с первыми клиентами"},
            {"label": "Q3 2026", "title": "Масштабирование", "body": "Выход на новые рынки"},
            {"label": "Q4 2026", "title": "Оптимизация", "body": "Снижение издержек"},
        ]},
        {"type": "content", "layout": "comparison", "title": "Сравнение подходов", "comparison": {
            "left": {"title": "Вариант A", "bullets": ["Быстрее", "Дешевле"]},
            "right": {"title": "Вариант B", "bullets": ["Надёжнее", "Масштабируемее"]},
        }},
        {"type": "content", "layout": "comparison", "title": "Неполные данные -> должен стать cards", "comparison": {
            "left": {"title": "Только один вариант"},
        }},
        {"type": "content", "layout": "team", "title": "Команда (4 -> сетка 2x2)", "team": [
            {"name": "Анна Иванова", "role": "CEO", "note": "10 лет в индустрии"},
            {"name": "Пётр Сидоров", "role": "CTO", "note": "Экс-Google"},
            {"name": "SpaceX (Starlink)", "role": "Партнёр", "note": "Инициалы не должны включать скобку"},
            {"name": "Олег Смирнов", "role": "CMO", "note": "Экс-стартапер"},
        ]},
        {"type": "content", "layout": "team", "title": "Команда (5 -> неполная строка центрируется)", "team": [
            {"name": "Игрок 1", "role": "Роль 1"},
            {"name": "Игрок 2", "role": "Роль 2"},
            {"name": "Игрок 3", "role": "Роль 3"},
            {"name": "Игрок 4", "role": "Роль 4"},
            {"name": "Игрок 5", "role": "Роль 5"},
        ]},
        {"type": "content", "layout": "big_photo", "title": "Вдохновляющий момент", "subtitle": "Минимум текста, максимум визуала", "image_query": "sunrise mountains"},
        {"type": "section", "title": "Разделитель без layout -> должен стать divider"},
        {"type": "section", "layout": "quote", "title": "Качество — это не случайность, а привычка.", "subtitle": "Аристотель"},
        {"type": "closing", "title": "Что делать дальше (буллеты)", "bullets": ["Утвердить бюджет", "Назначить ответственных", "Запустить пилот"]},
        {"type": "closing", "title": "Что делать дальше (body)", "body": "Связный абзац вместо списка следующих шагов."},
    ],
}


def _check_textbox_overflow(pptx_path: str) -> list[tuple[int, str]]:
    """Реальные (непустые) текстовые рамки, чьи границы вылезают за пределы
    слайда — верный признак будущего наложения/обрезки. Декоративные пустые
    фигуры (например, bleed-круги титульного слайда) не считаются, т.к. у
    них нет текста и вылезание за край — часть замысла."""
    from pptx import Presentation as _P

    prs = _P(pptx_path)
    issues = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if shape.left < 0 or shape.top < 0 or right > prs.slide_width or bottom > prs.slide_height:
                issues.append((i, shape.text_frame.text[:40]))
    return issues


async def main() -> None:
    ok = True
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        data = normalize_presentation(RAW, visual_format="balanced")

        path = await build_presentation(data, user_id=1, theme_preset="warm")
        print("OK no-bg:", path)
        issues = _check_textbox_overflow(path)
        if issues:
            ok = False
            print("  OVERFLOW on real text boxes:", issues)

        light_bg = tmp_path / "bg_light.jpg"
        Image.new("RGB", (1600, 1000), (235, 230, 220)).save(light_bg)
        path2 = await build_presentation(data, user_id=2, theme_preset="corporate", background_path=light_bg)
        print("OK light-bg:", path2)
        issues2 = _check_textbox_overflow(path2)
        if issues2:
            ok = False
            print("  OVERFLOW on real text boxes:", issues2)

        dark_bg = tmp_path / "bg_dark.jpg"
        Image.new("RGB", (1600, 1000), (15, 18, 25)).save(dark_bg)
        path3 = await build_presentation(data, user_id=3, theme_preset="dark", background_path=dark_bg)
        print("OK dark-bg:", path3)
        issues3 = _check_textbox_overflow(path3)
        if issues3:
            ok = False
            print("  OVERFLOW on real text boxes:", issues3)

        for fmt in ("image_only", "image_heavy", "balanced"):
            fmt_data = normalize_presentation(RAW, visual_format=fmt)
            bullets = fmt_data["slides"][1]["bullets"]
            print(f"{fmt}: bullets on slide 2 ->", bullets)

    print("\nPASS" if ok else "\nFAIL — see OVERFLOW lines above")
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
