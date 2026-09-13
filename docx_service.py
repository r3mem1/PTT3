"""Извлечение текста из загруженного пользователем .docx.

Сохраняем структуру заголовков (помечаем их "## ") — это даёт модели
подсказку об иерархии документа при построении выжимки, вместо плоского
текста без ориентиров.
"""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document

log = logging.getLogger(__name__)

# Ограничение на длину текста, уходящего в промпт: экономит токены и
# защищает от переполнения контекста на очень длинных документах.
MAX_CHARS = 12000


class DocxServiceError(Exception):
    """Ожидаемый сбой при чтении Word-документа."""


def extract_structured_text(path: Path) -> tuple[str, bool]:
    """Возвращает (текст с пометками заголовков, флаг "текст был обрезан")."""
    try:
        doc = Document(str(path))
    except Exception as e:  # python-docx кидает разные исключения на битые/чужие файлы
        log.warning("Не удалось прочитать .docx %s: %s", path, e)
        raise DocxServiceError(
            "Не удалось прочитать файл — убедитесь, что это корректный .docx (Word), а не .doc/.pdf/переименованный файл."
        ) from e

    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower() if para.style else ""
        if "heading" in style_name or "title" in style_name:
            lines.append(f"## {text}")
        else:
            lines.append(text)

    # Простые таблицы тоже часто несут важные тезисы (например, сравнение вариантов).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    full_text = "\n".join(lines).strip()
    if not full_text:
        raise DocxServiceError("В документе не нашлось текста для обработки.")

    truncated = len(full_text) > MAX_CHARS
    if truncated:
        full_text = full_text[:MAX_CHARS]
    return full_text, truncated
