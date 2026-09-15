"""
Собирает .pptx: раскладки рисуются шейпами поверх Blank-макета,
чтобы колода выглядела как единый дизайн, а не набор стандартных
плейсхолдеров PowerPoint.
"""

from __future__ import annotations

import logging
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Length, Pt

import image_service
from config import settings
from image_service import find_and_download_image
from models import PresentationData, SlideData
from theme import (
    BODY_SIZE,
    FOOTER_H,
    GUTTER,
    KICKER_SIZE,
    LOGO_SIZE,
    MARGIN,
    METRIC_SIZE,
    RAIL,
    TITLE_SIZE,
    get_theme,
)

log = logging.getLogger(__name__)

LAYOUT_BLANK = 6


@dataclass
class BuildContext:
    prs: Presentation
    deck_title: str
    index: int
    total: int
    used_images: list[Path]
    theme: dict = field(default_factory=lambda: get_theme(None))
    logo_path: Path | None = None
    background_path: Path | None = None
    bg_is_dark: bool = False


def _fill_solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _scrim(slide, prs, color: RGBColor, alpha: float) -> None:
    """Полупрозрачная плашка на весь слайд поверх фонового фото — гарантирует
    читаемость текста независимо от локального контраста конкретного кадра.

    python-pptx не даёт прозрачность заливки через публичный API, поэтому
    дописываем <a:alpha> в OOXML заливки напрямую — стандартный обходной путь."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _fill_solid(shape, color)
    srgb = shape.fill.fore_color._color._xClr
    alpha_el = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
    srgb.append(alpha_el)


def _strip_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        el = ph._element
        el.getparent().remove(el)


def _set_run(paragraph, text: str, *, font: str, size, color: RGBColor, bold=False, italic=False) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = size
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color


def _textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    font: str,
    size,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
    align=None,
    anchor: str = "t",
    shrink_to_fit: bool = False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if shrink_to_fit:
        # Длинный буллет/заголовок ужимается по шрифту вместо того, чтобы
        # вылезти за пределы карточки/колонки.
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.clear()
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    _set_run(p, text, font=font, size=size, color=color, bold=bold, italic=italic)
    tf._txBody.bodyPr.set("anchor", anchor)
    return box


def _kicker_text(slide_data: SlideData, fallback: str) -> str:
    return (slide_data.get("kicker") or fallback).upper()


def _add_logo(slide, ctx: BuildContext) -> None:
    """Небольшой логотип в правом нижнем углу, если пользователь его загрузил."""
    if ctx.logo_path is None or not Path(ctx.logo_path).is_file():
        return
    left = ctx.prs.slide_width - MARGIN - LOGO_SIZE
    top = ctx.prs.slide_height - FOOTER_H - LOGO_SIZE - Inches(0.08)
    try:
        with Image.open(str(ctx.logo_path)) as im:
            w, h = im.size
        ratio = w / h
        if ratio >= 1:
            box_w, box_h = LOGO_SIZE, LOGO_SIZE / ratio
        else:
            box_w, box_h = LOGO_SIZE * ratio, LOGO_SIZE
        pic = slide.shapes.add_picture(str(ctx.logo_path), left, top, width=box_w, height=box_h)
        pic.line.fill.background()
    except Exception:
        log.debug("Не удалось наложить логотип %s", ctx.logo_path)


def _add_footer(slide, ctx: BuildContext, *, on_dark: bool = False) -> None:
    theme = ctx.theme
    color = theme["on_dark_muted"] if on_dark else theme["muted_color"]
    top = ctx.prs.slide_height - FOOTER_H
    width = ctx.prs.slide_width - MARGIN * 2
    label = ctx.deck_title[:42]
    _textbox(
        slide, MARGIN, top, width * 0.75, FOOTER_H, label,
        font=theme["body_font"], size=Pt(10), color=color,
    )
    _textbox(
        slide, MARGIN, top, width, FOOTER_H, f"{ctx.index:02d}  /  {ctx.total:02d}",
        font=theme["body_font"], size=Pt(10), color=color, align=PP_ALIGN.RIGHT,
    )
    _add_logo(slide, ctx)


def _add_rail(slide, prs, color: RGBColor) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, RAIL, prs.slide_height)
    _fill_solid(bar, color)


def _picture(slide, image_path: Path, left, top, width, height):
    path = str(image_path)
    with Image.open(path) as im:
        img_w, img_h = im.size
    img_ratio = img_w / img_h
    box_ratio = width / height
    pic = slide.shapes.add_picture(path, left, top, width=width, height=height)
    if img_ratio > box_ratio:
        crop = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    pic.line.fill.background()
    return pic


def _card_shape(slide, left, top, width, height, theme: dict):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill_solid(shape, theme["card"])
    try:
        shape.adjustments[0] = 0.06
    except (AttributeError, IndexError):
        pass
    shape.line.color.rgb = theme["accent_soft"]
    shape.line.width = Pt(1)
    return shape


def _decorative_panel(slide, prs, left, top, width, height, theme: dict) -> None:
    """Заполняет правую панель титульного слайда геометрией, если картинки нет.

    Пустая заливка выглядит недоделанной — несколько полупрозрачно-имитирующих
    кругов акцентным/мягким цветом дают ощущение осознанного дизайна.
    """
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill_solid(base, theme["bg_dark"])

    big = int(min(width, height) * 0.9)
    circle1 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + width - big // 2, top - big // 4, big, big
    )
    _fill_solid(circle1, theme["panel"])

    mid = int(min(width, height) * 0.55)
    circle2 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + int(width * 0.1), top + height - mid // 2, mid, mid
    )
    _fill_solid(circle2, theme["accent"])

    small = int(min(width, height) * 0.22)
    circle3 = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        left + int(width * 0.55),
        top + int(height * 0.35),
        small,
        small,
    )
    _fill_solid(circle3, theme["accent_soft"])


async def _maybe_image(slide_data: SlideData, ctx: BuildContext) -> Path | None:
    query = slide_data.get("image_query")
    if not query:
        return None
    path = await find_and_download_image(query)
    if path is not None:
        ctx.used_images.append(Path(path))
    return path


def _new_slide(ctx: BuildContext, *, dark: bool = False) -> tuple:
    """Создаёт пустой слайд и красит фон.

    Возвращает (slide, on_dark) — on_dark говорит вызывающей функции, каким
    набором цветов текста пользоваться. Обычно on_dark == dark, но если
    пользователь загрузил кастомный фон колоды, реальный фон — это фото, и
    его яркость (ctx.bg_is_dark) может переопределить выбор цвета текста
    даже для контентных слайдов, которые по умолчанию светлые.
    """
    slide = ctx.prs.slides.add_slide(ctx.prs.slide_layouts[LAYOUT_BLANK])
    _strip_placeholders(slide)

    if ctx.background_path is not None:
        _picture(slide, ctx.background_path, 0, 0, ctx.prs.slide_width, ctx.prs.slide_height)
        on_dark = dark or ctx.bg_is_dark
        scrim_color = ctx.theme["bg_dark"] if on_dark else ctx.theme["bg"]
        _scrim(slide, ctx.prs, scrim_color, alpha=0.55 if dark else 0.45)
        return slide, on_dark

    _slide_bg(slide, ctx.theme["bg_dark"] if dark else ctx.theme["bg"])
    return slide, dark


def _add_title_block(
    slide, prs, slide_data: SlideData, *, kicker: str, theme: dict, on_dark: bool = False
) -> Length:
    title_color = theme["on_dark"] if on_dark else theme["title_color"]
    kicker_color = theme["accent"] if not on_dark else theme["accent_soft"]
    left = MARGIN + Inches(0.12)
    width = prs.slide_width - MARGIN * 2
    _textbox(
        slide, left, Inches(0.38), width, Inches(0.32), kicker,
        font=theme["body_font"], size=KICKER_SIZE, color=kicker_color, bold=True,
    )
    _textbox(
        slide, left, Inches(0.68), width, Inches(1.15), slide_data.get("title") or "",
        font=theme["heading_font"], size=TITLE_SIZE, color=title_color, bold=True,
        shrink_to_fit=True,
    )
    return Inches(1.95)


async def _add_title_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, _ = _new_slide(ctx, dark=True)
    prs = ctx.prs

    if ctx.background_path is not None:
        # Кастомный фон колоды уже занимает весь слайд — вторая, "своя"
        # картинка титульного слайда только конфликтовала бы с ним, поэтому
        # пропускаем панель/картинку и рисуем текст поверх общего фона.
        _add_rail(slide, prs, theme["accent"])
        pad = MARGIN + Inches(0.4)
        text_w = prs.slide_width - pad - MARGIN
        _textbox(
            slide, pad, Inches(2.1), text_w, Inches(0.35),
            _kicker_text(slide_data, "Briefing"),
            font=theme["body_font"], size=KICKER_SIZE, color=theme["accent_soft"], bold=True,
        )
        _textbox(
            slide, pad, Inches(2.55), text_w, Inches(2.4),
            slide_data.get("title") or ctx.deck_title,
            font=theme["heading_font"], size=Pt(40), color=theme["on_dark"], bold=True,
            shrink_to_fit=True,
        )
        subtitle = slide_data.get("subtitle") or ""
        if subtitle:
            _textbox(
                slide, pad, Inches(5.1), text_w, Inches(1.0), subtitle,
                font=theme["body_font"], size=Pt(16), color=theme["on_dark_muted"],
            )
        _add_footer(slide, ctx, on_dark=True)
        return

    panel_w = int(prs.slide_width * 0.56)
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, panel_w, prs.slide_height)
    _fill_solid(panel, theme["panel"])
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel_w, 0, RAIL, prs.slide_height)
    _fill_solid(accent, theme["accent"])

    image_path = await _maybe_image(slide_data, ctx)
    right_w = prs.slide_width - panel_w
    if image_path is not None:
        _picture(slide, image_path, panel_w, 0, right_w, prs.slide_height)
    else:
        # Раньше здесь оставалось пустое цветное поле — теперь декоративная геометрия.
        _decorative_panel(slide, prs, panel_w, 0, right_w, prs.slide_height, theme)

    pad = Inches(0.7)
    _textbox(
        slide, pad, Inches(1.55), panel_w - pad * 1.4, Inches(0.35),
        _kicker_text(slide_data, "Briefing"),
        font=theme["body_font"], size=KICKER_SIZE, color=theme["accent_soft"], bold=True,
    )
    _textbox(
        slide, pad, Inches(2.0), panel_w - pad * 1.4, Inches(2.4),
        slide_data.get("title") or ctx.deck_title,
        font=theme["heading_font"], size=Pt(40), color=theme["on_dark"], bold=True,
        shrink_to_fit=True,
    )
    subtitle = slide_data.get("subtitle") or ""
    if subtitle:
        _textbox(
            slide, pad, Inches(4.6), panel_w - pad * 1.4, Inches(1.4), subtitle,
            font=theme["body_font"], size=Pt(16), color=theme["on_dark_muted"],
        )
    _add_footer(slide, ctx, on_dark=True)


def _add_points(slide, left, top, width, bullets: list[str], theme: dict, *, light: bool = True) -> None:
    row_h = Inches(0.92)
    num_color = theme["accent"]
    text_color = theme["body_color"] if light else theme["on_dark"]
    for i, bullet in enumerate(bullets[:4]):
        y = top + row_h * i
        _textbox(
            slide, left, y, Inches(0.55), Inches(0.4), f"{i + 1:02d}",
            font=theme["heading_font"], size=Pt(16), color=num_color, bold=True,
        )
        _textbox(
            slide, left + Inches(0.62), y, width - Inches(0.62), Inches(0.82), bullet,
            font=theme["body_font"], size=BODY_SIZE, color=text_color, shrink_to_fit=True,
        )


def _add_text_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Key points"), theme=theme, on_dark=on_dark
    )
    bullets = slide_data.get("bullets") or []
    if bullets:
        _add_points(
            slide, MARGIN + Inches(0.12), content_top, ctx.prs.slide_width - MARGIN * 2, bullets, theme,
            light=not on_dark,
        )
    _add_footer(slide, ctx, on_dark=on_dark)


async def _add_image_text_slide(ctx: BuildContext, slide_data: SlideData, image_on_left: bool) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Insight"), theme=theme, on_dark=on_dark
    )

    gap = GUTTER
    body_top = content_top
    body_h = ctx.prs.slide_height - body_top - FOOTER_H - Inches(0.15)
    col_w = (ctx.prs.slide_width - MARGIN * 2 - gap) * 0.46
    text_w = ctx.prs.slide_width - MARGIN * 2 - gap - col_w
    img_left = MARGIN if image_on_left else MARGIN + text_w + gap
    text_left = MARGIN + col_w + gap if image_on_left else MARGIN

    image_path = await _maybe_image(slide_data, ctx)
    if image_path is not None:
        pic_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_left, body_top, col_w, body_h)
        _fill_solid(pic_box, theme["accent_soft"])
        _picture(slide, image_path, img_left, body_top, col_w, body_h)
    else:
        text_left = MARGIN
        text_w = ctx.prs.slide_width - MARGIN * 2

    bullets = slide_data.get("bullets") or []
    if bullets:
        _add_points(slide, text_left, body_top, text_w, bullets, theme, light=not on_dark)
    _add_footer(slide, ctx, on_dark=on_dark)


def _add_cards_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Framework"), theme=theme, on_dark=on_dark
    )
    cards = slide_data.get("cards") or []
    if not cards:
        bullets = slide_data.get("bullets") or []
        cards = [{"title": b, "body": ""} for b in bullets[:3]]

    n = max(1, min(3, len(cards)))
    available = ctx.prs.slide_width - MARGIN * 2
    card_w = (available - GUTTER * (n - 1)) / n
    height = ctx.prs.slide_height - content_top - FOOTER_H - Inches(0.2)
    for i, card in enumerate(cards[:n]):
        left = MARGIN + i * (card_w + GUTTER)
        _card_shape(slide, left, content_top, card_w, height, theme)
        inner = left + Inches(0.28)
        inner_w = card_w - Inches(0.56)
        _textbox(
            slide, inner, content_top + Inches(0.28), inner_w, Inches(0.36), f"{i + 1:02d}",
            font=theme["heading_font"], size=Pt(14), color=theme["accent"], bold=True,
        )
        _textbox(
            slide, inner, content_top + Inches(0.7), inner_w, Inches(1.3), card.get("title") or "",
            font=theme["heading_font"], size=Pt(20), color=theme["title_color"], bold=True,
            shrink_to_fit=True,
        )
        _textbox(
            slide, inner, content_top + Inches(2.1), inner_w, height - Inches(2.4), card.get("body") or "",
            font=theme["body_font"], size=Pt(14), color=theme["body_color"], shrink_to_fit=True,
        )
    _add_footer(slide, ctx, on_dark=on_dark)


def _add_stat_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Signal"), theme=theme, on_dark=on_dark
    )
    metric = slide_data.get("metric") or (slide_data.get("bullets") or ["—"])[0]
    label = slide_data.get("metric_label") or slide_data.get("subtitle") or ""
    label_color = theme["on_dark"] if on_dark else theme["title_color"]
    body_color = theme["on_dark_muted"] if on_dark else theme["body_color"]
    _textbox(
        slide, MARGIN, Inches(2.15), ctx.prs.slide_width - MARGIN * 2, Inches(1.6), metric,
        font=theme["heading_font"], size=METRIC_SIZE, color=theme["accent"], bold=True,
        shrink_to_fit=True,
    )
    if label:
        _textbox(
            slide, MARGIN, Inches(3.75), ctx.prs.slide_width - MARGIN * 2, Inches(0.7), label,
            font=theme["heading_font"], size=Pt(22), color=label_color, bold=True,
            shrink_to_fit=True,
        )
    bullets = slide_data.get("bullets") or []
    if bullets:
        _textbox(
            slide, MARGIN, Inches(4.55), ctx.prs.slide_width - MARGIN * 2, Inches(1.6), bullets[0],
            font=theme["body_font"], size=Pt(16), color=body_color, shrink_to_fit=True,
        )
    _add_footer(slide, ctx, on_dark=on_dark)


def _add_chart_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Data"), theme=theme, on_dark=on_dark
    )

    raw = slide_data.get("chart")
    if not raw:
        _add_footer(slide, ctx, on_dark=on_dark)
        return

    chart_type_map = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "pie": XL_CHART_TYPE.PIE,
        "line": XL_CHART_TYPE.LINE_MARKERS,
    }
    xl_type = chart_type_map.get(raw["type"], XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_data = CategoryChartData()
    chart_data.categories = raw["categories"]
    for series in raw["series"]:
        chart_data.add_series(series["name"], series["values"])

    width = ctx.prs.slide_width - MARGIN * 2
    height = ctx.prs.slide_height - content_top - FOOTER_H - Inches(0.2)
    graphic_frame = slide.shapes.add_chart(xl_type, MARGIN, content_top, width, height, chart_data)
    chart = graphic_frame.chart

    text_color = theme["on_dark"] if on_dark else theme["body_color"]
    series_colors = [theme["accent"], theme["accent_soft"], theme["muted_color"]]
    for i, plot_series in enumerate(chart.plots[0].series):
        plot_series.format.fill.solid()
        plot_series.format.fill.fore_color.rgb = series_colors[i % len(series_colors)]

    chart.has_legend = len(raw["series"]) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    if xl_type == XL_CHART_TYPE.PIE:
        plot = chart.plots[0]
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.number_format = "0%"
        data_labels.number_format_is_linked = False
        data_labels.font.size = Pt(12)
        data_labels.font.color.rgb = text_color
        # Сегменты pie красим по очереди — иначе все получат цвет одной серии.
        palette = [theme["accent"], theme["accent_soft"], theme["muted_color"], theme["panel"]]
        for i, point in enumerate(chart.plots[0].series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = palette[i % len(palette)]
    else:
        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(11)
        category_axis.tick_labels.font.color.rgb = text_color
        category_axis.format.line.color.rgb = theme["muted_color"]
        value_axis = chart.value_axis
        value_axis.tick_labels.font.size = Pt(11)
        value_axis.tick_labels.font.color.rgb = text_color
        value_axis.format.line.fill.background()
        value_axis.has_major_gridlines = False

    _add_footer(slide, ctx, on_dark=on_dark)


def _add_timeline_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Roadmap"), theme=theme, on_dark=on_dark
    )
    steps = (slide_data.get("timeline") or [])[:5]
    if not steps:
        _add_footer(slide, ctx, on_dark=on_dark)
        return

    title_color = theme["on_dark"] if on_dark else theme["title_color"]
    body_color = theme["on_dark_muted"] if on_dark else theme["body_color"]

    left0 = MARGIN + Inches(0.12)
    available = ctx.prs.slide_width - MARGIN * 2 - Inches(0.12)
    col_w = available // len(steps)
    marker_d = Inches(0.36)

    for i, step in enumerate(steps):
        left = left0 + col_w * i
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, content_top, marker_d, marker_d)
        _fill_solid(marker, theme["accent"])
        _textbox(
            slide, left, content_top, marker_d, marker_d, str(i + 1),
            font=theme["heading_font"], size=Pt(14), color=theme["on_dark"], bold=True,
            align=PP_ALIGN.CENTER, anchor="ctr",
        )
        _textbox(
            slide, left, content_top + marker_d + Inches(0.15), col_w - Inches(0.2), Inches(0.3),
            step.get("label") or "", font=theme["body_font"], size=Pt(11), color=theme["accent"], bold=True,
        )
        _textbox(
            slide, left, content_top + marker_d + Inches(0.5), col_w - Inches(0.2), Inches(0.55),
            step.get("title") or "", font=theme["heading_font"], size=Pt(15), color=title_color, bold=True,
            shrink_to_fit=True,
        )
        body = step.get("body") or ""
        if body:
            _textbox(
                slide, left, content_top + marker_d + Inches(1.05), col_w - Inches(0.2), Inches(1.3),
                body, font=theme["body_font"], size=Pt(12), color=body_color, shrink_to_fit=True,
            )
    _add_footer(slide, ctx, on_dark=on_dark)


def _add_comparison_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Trade-off"), theme=theme, on_dark=on_dark
    )
    comparison = slide_data.get("comparison") or {}
    left_side = comparison.get("left") or {"title": "Вариант A", "bullets": []}
    right_side = comparison.get("right") or {"title": "Вариант B", "bullets": []}

    available = ctx.prs.slide_width - MARGIN * 2
    height = ctx.prs.slide_height - content_top - FOOTER_H - Inches(0.2)
    col_w = int((available - GUTTER) / 2)

    for i, side in enumerate((left_side, right_side)):
        left = MARGIN + i * (col_w + GUTTER)
        _card_shape(slide, left, content_top, col_w, height, theme)
        inner = left + Inches(0.28)
        inner_w = col_w - Inches(0.56)
        _textbox(
            slide, inner, content_top + Inches(0.28), inner_w, Inches(0.5), side.get("title") or "",
            font=theme["heading_font"], size=Pt(18), color=theme["title_color"], bold=True,
            shrink_to_fit=True,
        )
        bullets = side.get("bullets") or []
        if bullets:
            _add_points(slide, inner, content_top + Inches(0.9), inner_w, bullets, theme)
    _add_footer(slide, ctx, on_dark=on_dark)


def _add_team_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, on_dark = _new_slide(ctx)
    _add_rail(slide, ctx.prs, theme["accent"])
    content_top = _add_title_block(
        slide, ctx.prs, slide_data, kicker=_kicker_text(slide_data, "Team"), theme=theme, on_dark=on_dark
    )
    members = (slide_data.get("team") or [])[:6]
    if not members:
        _add_footer(slide, ctx, on_dark=on_dark)
        return

    cols = min(3, len(members))
    rows = (len(members) + cols - 1) // cols
    available_w = ctx.prs.slide_width - MARGIN * 2
    available_h = ctx.prs.slide_height - content_top - FOOTER_H - Inches(0.2)
    cell_w = available_w // cols
    cell_h = available_h // rows
    avatar_d = Inches(0.85)

    title_color = theme["on_dark"] if on_dark else theme["title_color"]
    body_color = theme["on_dark_muted"] if on_dark else theme["body_color"]

    for i, member in enumerate(members):
        row, col = divmod(i, cols)
        cell_left = MARGIN + cell_w * col
        cell_top = content_top + cell_h * row
        avatar_left = cell_left + (cell_w - avatar_d) // 2
        avatar = slide.shapes.add_shape(MSO_SHAPE.OVAL, avatar_left, cell_top, avatar_d, avatar_d)
        _fill_solid(avatar, theme["accent_soft"])
        initials = "".join(part[0].upper() for part in (member.get("name") or "?").split()[:2]) or "?"
        _textbox(
            slide, avatar_left, cell_top, avatar_d, avatar_d, initials,
            font=theme["heading_font"], size=Pt(22), color=theme["accent"], bold=True,
            align=PP_ALIGN.CENTER, anchor="ctr",
        )
        text_top = cell_top + avatar_d + Inches(0.12)
        _textbox(
            slide, cell_left, text_top, cell_w, Inches(0.35), member.get("name") or "",
            font=theme["heading_font"], size=Pt(15), color=title_color, bold=True,
            align=PP_ALIGN.CENTER, shrink_to_fit=True,
        )
        role = member.get("role") or ""
        if role:
            _textbox(
                slide, cell_left, text_top + Inches(0.34), cell_w, Inches(0.3), role,
                font=theme["body_font"], size=Pt(12), color=theme["accent"], align=PP_ALIGN.CENTER, bold=True,
            )
        note = member.get("note") or ""
        if note:
            _textbox(
                slide, cell_left, text_top + Inches(0.62), cell_w, Inches(0.6), note,
                font=theme["body_font"], size=Pt(11), color=body_color, align=PP_ALIGN.CENTER,
                shrink_to_fit=True,
            )
    _add_footer(slide, ctx, on_dark=on_dark)


async def _add_big_photo_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, _ = _new_slide(ctx, dark=True)

    if ctx.background_path is None:
        # Без кастомного фона колоды — своя полноэкранная картинка по теме слайда.
        image_path = await _maybe_image(slide_data, ctx)
        if image_path is not None:
            _picture(slide, image_path, 0, 0, ctx.prs.slide_width, ctx.prs.slide_height)
        _scrim(slide, ctx.prs, theme["bg_dark"], alpha=0.55)
    # Если кастомный фон уже задан, _new_slide уже нарисовал его и скрим —
    # рисовать поверх ещё одну картинку означало бы просто перекрыть фон.

    pad = Inches(0.9)
    text_w = ctx.prs.slide_width - pad * 2
    _textbox(
        slide, pad, ctx.prs.slide_height - Inches(2.6), text_w, Inches(1.5),
        slide_data.get("title") or "", font=theme["heading_font"], size=Pt(34), color=theme["on_dark"],
        bold=True, shrink_to_fit=True,
    )
    subtitle = slide_data.get("subtitle") or ""
    if subtitle:
        _textbox(
            slide, pad, ctx.prs.slide_height - Inches(1.15), text_w, Inches(0.7), subtitle,
            font=theme["body_font"], size=Pt(16), color=theme["on_dark_muted"],
        )
    _add_footer(slide, ctx, on_dark=True)


def _add_divider_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, _ = _new_slide(ctx, dark=True)
    _add_rail(slide, ctx.prs, theme["accent"])
    _textbox(
        slide, MARGIN, Inches(2.3), ctx.prs.slide_width - MARGIN * 2, Inches(0.4),
        _kicker_text(slide_data, "Section"),
        font=theme["body_font"], size=KICKER_SIZE, color=theme["accent_soft"], bold=True,
    )
    _textbox(
        slide, MARGIN, Inches(2.75), ctx.prs.slide_width - MARGIN * 2, Inches(2.2),
        slide_data.get("title") or "",
        font=theme["heading_font"], size=Pt(40), color=theme["on_dark"], bold=True,
        shrink_to_fit=True,
    )
    subtitle = slide_data.get("subtitle") or ""
    if subtitle:
        _textbox(
            slide, MARGIN, Inches(4.85), ctx.prs.slide_width - MARGIN * 2, Inches(0.8), subtitle,
            font=theme["body_font"], size=Pt(16), color=theme["on_dark_muted"],
        )
    _add_footer(slide, ctx, on_dark=True)


def _add_quote_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, _ = _new_slide(ctx, dark=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, ctx.prs.slide_width, RAIL)
    _fill_solid(bar, theme["accent"])
    quote = slide_data.get("title") or ""
    _textbox(
        slide, Inches(1.1), Inches(2.1), ctx.prs.slide_width - Inches(2.2), Inches(2.8),
        f"“{quote}”" if quote and not quote.startswith("“") else quote,
        font=theme["heading_font"], size=Pt(32), color=theme["on_dark"], bold=True,
        shrink_to_fit=True,
    )
    attribution = slide_data.get("subtitle") or slide_data.get("kicker") or ""
    if attribution:
        _textbox(
            slide, Inches(1.1), Inches(5.1), ctx.prs.slide_width - Inches(2.2), Inches(0.6),
            attribution,
            font=theme["body_font"], size=Pt(16), color=theme["accent_soft"], italic=True,
        )
    _add_footer(slide, ctx, on_dark=True)


def _add_closing_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    theme = ctx.theme
    slide, _ = _new_slide(ctx, dark=True)
    _add_rail(slide, ctx.prs, theme["accent"])
    _textbox(
        slide, MARGIN, Inches(0.55), ctx.prs.slide_width - MARGIN * 2, Inches(0.35),
        _kicker_text(slide_data, "Next move"),
        font=theme["body_font"], size=KICKER_SIZE, color=theme["accent_soft"], bold=True,
    )
    _textbox(
        slide, MARGIN, Inches(1.0), ctx.prs.slide_width - MARGIN * 2, Inches(1.6),
        slide_data.get("title") or "Что делать дальше",
        font=theme["heading_font"], size=Pt(36), color=theme["on_dark"], bold=True,
        shrink_to_fit=True,
    )
    bullets = slide_data.get("bullets") or []
    if bullets:
        _add_points(slide, MARGIN, Inches(2.8), ctx.prs.slide_width - MARGIN * 2, bullets, theme, light=False)
    _add_footer(slide, ctx, on_dark=True)


async def _dispatch_slide(ctx: BuildContext, slide_data: SlideData) -> None:
    slide_type = slide_data.get("type", "content")
    layout = slide_data.get("layout", "text")

    if slide_type == "title":
        await _add_title_slide(ctx, slide_data)
        return
    if slide_type == "closing":
        _add_closing_slide(ctx, slide_data)
        return
    if slide_type == "section":
        if layout == "quote":
            _add_quote_slide(ctx, slide_data)
        else:
            _add_divider_slide(ctx, slide_data)
        return
    if layout == "cards":
        _add_cards_slide(ctx, slide_data)
        return
    if layout == "stat":
        _add_stat_slide(ctx, slide_data)
        return
    if layout == "chart":
        _add_chart_slide(ctx, slide_data)
        return
    if layout == "timeline":
        _add_timeline_slide(ctx, slide_data)
        return
    if layout == "comparison":
        _add_comparison_slide(ctx, slide_data)
        return
    if layout == "team":
        _add_team_slide(ctx, slide_data)
        return
    if layout == "big_photo":
        await _add_big_photo_slide(ctx, slide_data)
        return
    if layout == "image_left":
        await _add_image_text_slide(ctx, slide_data, image_on_left=True)
        return
    if layout == "image_right":
        await _add_image_text_slide(ctx, slide_data, image_on_left=False)
        return
    _add_text_slide(ctx, slide_data)


def _safe_filename(title: str) -> str:
    cleaned = "".join(c for c in title if c.isalnum() or c in " _-").strip()
    return cleaned[:40] or "presentation"


async def build_presentation(
    data: PresentationData,
    user_id: int,
    *,
    theme_preset: str | None = None,
    logo_path: Path | None = None,
    background_path: Path | None = None,
) -> str:
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(settings.template_path))
    used_images: list[Path] = []
    slides = data.get("slides") or []
    if not slides:
        raise ValueError("В ответе ИИ нет ни одного слайда.")

    resolved_background: Path | None = None
    bg_is_dark = False
    if background_path is not None and Path(background_path).is_file():
        resolved_background = Path(background_path)
        try:
            bg_is_dark = image_service.analyze_brightness(resolved_background) < 128
        except Exception:
            log.debug("Не удалось проанализировать яркость фона %s", resolved_background)
            resolved_background = None

    ctx = BuildContext(
        prs=prs,
        deck_title=data.get("title") or "presentation",
        index=1,
        total=len(slides),
        used_images=used_images,
        theme=get_theme(theme_preset),
        logo_path=logo_path,
        background_path=resolved_background,
        bg_is_dark=bg_is_dark,
    )

    try:
        for i, slide_data in enumerate(slides, start=1):
            ctx.index = i
            await _dispatch_slide(ctx, slide_data)

        file_name = f"{user_id}_{_safe_filename(ctx.deck_title)}_{uuid.uuid4().hex[:6]}.pptx"
        file_path = settings.output_dir / file_name
        prs.save(str(file_path))
        log.info("Сохранён файл %s (%s слайдов)", file_path, len(slides))
        return str(file_path)
    finally:
        for image_path in used_images:
            try:
                image_path.unlink(missing_ok=True)
            except OSError:
                log.debug("Не удалось удалить временную картинку %s", image_path)
